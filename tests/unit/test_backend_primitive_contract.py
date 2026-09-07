"""Contract tests for the two backends' primitive twins (Spec sec 4.0 rule 3;
docs/backend_eligibility_build_plan.md W5/W6).

The whole build rests on one promise: **an action produces the same result whichever backend
it lands on**, because the workflow author never chose the backend and cannot see it. That
promise is only credible if it is asserted by *one* test body run against *both* backends —
two separate test files would drift apart precisely where they matter, and each would still
pass.

So every test here is written once and parameterised over `("file", "xlw")`. The `xlw` half is
skipped without a live Excel; the `file` half runs everywhere. No mocks — a mocked backend
would prove nothing about whether the twins actually agree.

Two divergences are deliberate and are *not* bugs:

* xlwings returns floats (`5.0`) where openpyxl returns ints (`5`), because Excel has one
  numeric type. Assertions here use `==` and never assert on type.
* `"autofit"` genuinely differs: openpyxl has no rendering engine, so the file backend
  approximates by string length while Excel measures the rendered glyphs. Only "a sensible
  positive width was set" is contractual.
"""

from collections.abc import Iterator
from pathlib import Path
from typing import Any, NamedTuple

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from excel_runner import backends
from tests.unit.conftest import requires_excel


class Opened(NamedTuple):
    """One workbook, open on one backend, with that backend's primitive table."""

    backend: str
    primitives: backends.BackendPrimitives
    handle: Any


def _fixture_workbook(path: Path) -> Path:
    """A small workbook with headers, data, a second sheet and a defined name — enough for
    every primitive to have something real to act on."""
    workbook = openpyxl.Workbook()
    data = workbook.active
    assert data is not None
    data.title = "Data"
    data["A2"] = "Region"
    data["B2"] = "Amount"
    data["C2"] = "Notes"
    rows = [("North", 10, "first"), ("South", 20, "second"), ("East", 30, "third")]
    for offset, (region, amount, note) in enumerate(rows):
        data.cell(row=3 + offset, column=1, value=region)
        data.cell(row=3 + offset, column=2, value=amount)
        data.cell(row=3 + offset, column=3, value=note)
    workbook.create_sheet("Notes")["A1"] = "aside"
    workbook.defined_names.add(DefinedName("SouthAmount", attr_text="Data!$B$4"))
    workbook.save(path)
    return path


@pytest.fixture(
    params=[
        pytest.param("file", id="file"),
        pytest.param("xlw", id="xlw", marks=requires_excel),
    ]
)
def opened(request: pytest.FixtureRequest, tmp_path: Path) -> Iterator[Opened]:
    """The same fixture workbook, open read-write on whichever backend is under test."""
    backend: str = request.param
    path = _fixture_workbook(tmp_path / "contract.xlsx")
    primitives = backends.primitives(backend)  # type: ignore[arg-type]
    if backend == "file":
        yield Opened(
            backend, primitives, backends.open_workbook(str(path), mode="read_write")
        )
        return
    registry = backends.OwnedInstanceRegistry()
    try:
        app = registry.spawn()
        handle = backends.xlw_open_workbook(app, str(path), mode="read_write")
        yield Opened(backend, primitives, handle)
    finally:
        registry.close_owned()


class TestReadContract:
    def test_resolve_sheet_names_lists_every_sheet_in_workbook_order(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.resolve_sheet_names(opened.handle, "all") == [
            "Data",
            "Notes",
        ]

    def test_resolve_sheet_names_passes_an_explicit_list_through(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.resolve_sheet_names(opened.handle, ["Notes"]) == [
            "Notes"
        ]

    def test_resolve_sheet_names_filters_by_regex(self, opened: Opened) -> None:
        assert opened.primitives.resolve_sheet_names(
            opened.handle, {"matching": "^Not"}
        ) == ["Notes"]

    def test_resolve_range_expands_a_defined_name_to_its_own_sheet_and_cell(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.resolve_range(
            opened.handle, "Notes", "SouthAmount"
        ) == (
            "Data",
            "B4",
        )

    def test_resolve_range_passes_plain_a1_notation_through(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.resolve_range(opened.handle, "Data", "B3") == (
            "Data",
            "B3",
        )

    def test_read_range_returns_a_scalar_for_a_single_cell(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.read_range(opened.handle, "Data", "B3") == 10

    def test_read_range_returns_a_2d_block_for_a_range(self, opened: Opened) -> None:
        assert opened.primitives.read_range(opened.handle, "Data", "A3:B4") == [
            ["North", 10],
            ["South", 20],
        ]

    def test_read_range_follows_a_defined_name_to_another_sheet(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.read_range(opened.handle, "Notes", "SouthAmount") == 20

    def test_read_cells_maps_each_reference_to_its_value(self, opened: Opened) -> None:
        assert opened.primitives.read_cells(
            opened.handle, "Data", ["A3", "B5", "SouthAmount"]
        ) == {"A3": "North", "B5": 30, "SouthAmount": 20}

    def test_read_properties_returns_a_mapping_of_document_properties(
        self, opened: Opened
    ) -> None:
        """Only the shape is contractual — the values are Excel's/openpyxl's own defaults and
        legitimately differ between them (creator, timestamps)."""
        properties = opened.primitives.read_properties(opened.handle)

        assert isinstance(properties, dict)
        assert all(isinstance(key, str) for key in properties)


class TestWriteContract:
    def test_write_cell_is_visible_to_a_subsequent_read(self, opened: Opened) -> None:
        opened.primitives.write_cell(opened.handle, "Data", "B3", 99)

        assert opened.primitives.read_range(opened.handle, "Data", "B3") == 99

    def test_write_range_anchors_on_the_top_left_cell_only(
        self, opened: Opened
    ) -> None:
        """The block's size comes from `values`, not from the extent of `range` (PRD sec 11
        item 8) — so a one-cell `range` still writes the whole block."""
        opened.primitives.write_range(
            opened.handle, "Data", "A3", [["West", 40], ["Central", 50]]
        )

        assert opened.primitives.read_range(opened.handle, "Data", "A3:B4") == [
            ["West", 40],
            ["Central", 50],
        ]

    def test_set_column_width_applies_an_explicit_width(self, opened: Opened) -> None:
        opened.primitives.set_column_width(opened.handle, "Data", "B", 25)

        # Read back through each backend's own accessor — the *stored* width is the contract,
        # and there is no primitive that reads it.
        if opened.backend == "file":
            assert opened.handle["Data"].column_dimensions["B"].width == 25
        else:
            assert opened.handle.sheets["Data"].range("B:B").column_width == 25

    def test_set_column_width_autofit_produces_a_positive_width(
        self, opened: Opened
    ) -> None:
        """The exact number legitimately differs: openpyxl has no rendering engine and
        approximates by string length, Excel measures real glyphs. Only "something sensible
        was set" is contractual."""
        opened.primitives.set_column_width(opened.handle, "Data", "C", "autofit")

        if opened.backend == "file":
            width = opened.handle["Data"].column_dimensions["C"].width
        else:
            width = opened.handle.sheets["Data"].range("C:C").column_width
        assert width > 0


class TestSheetStructureContract:
    def test_create_sheet_appends_by_default(self, opened: Opened) -> None:
        opened.primitives.create_sheet(opened.handle, "Extra")

        assert opened.primitives.resolve_sheet_names(opened.handle, "all") == [
            "Data",
            "Notes",
            "Extra",
        ]

    def test_create_sheet_honours_an_explicit_index(self, opened: Opened) -> None:
        opened.primitives.create_sheet(opened.handle, "First", index=0)

        assert opened.primitives.resolve_sheet_names(opened.handle, "all") == [
            "First",
            "Data",
            "Notes",
        ]

    def test_create_sheet_rejects_a_duplicate_name(self, opened: Opened) -> None:
        with pytest.raises(ValueError):
            opened.primitives.create_sheet(opened.handle, "Data")

    def test_rename_sheet_changes_the_name_in_place(self, opened: Opened) -> None:
        opened.primitives.rename_sheet(opened.handle, "Notes", "Asides")

        assert opened.primitives.resolve_sheet_names(opened.handle, "all") == [
            "Data",
            "Asides",
        ]

    def test_delete_sheet_removes_it(self, opened: Opened) -> None:
        opened.primitives.delete_sheet(opened.handle, "Notes")

        assert opened.primitives.resolve_sheet_names(opened.handle, "all") == ["Data"]

    def test_delete_sheet_refuses_to_empty_the_workbook(self, opened: Opened) -> None:
        opened.primitives.delete_sheet(opened.handle, "Notes")

        with pytest.raises(ValueError):
            opened.primitives.delete_sheet(opened.handle, "Data")


class TestInsertRangeContract:
    def test_whole_column_insert_shifts_existing_content_right(
        self, opened: Opened
    ) -> None:
        opened.primitives.insert_range(opened.handle, "Data", "A:A")

        assert opened.primitives.read_range(opened.handle, "Data", "B3") == "North"

    def test_whole_row_insert_shifts_existing_content_down(
        self, opened: Opened
    ) -> None:
        opened.primitives.insert_range(opened.handle, "Data", "3:3")

        assert opened.primitives.read_range(opened.handle, "Data", "A4") == "North"

    def test_partial_range_insert_is_refused_identically_by_both_backends(
        self, opened: Opened
    ) -> None:
        """Same limitation, same message — a user must not be able to tell which backend they
        landed on from the error they get (PRD sec 11 item 12)."""
        with pytest.raises(NotImplementedError) as excinfo:
            opened.primitives.insert_range(opened.handle, "Data", "B2:C4")

        assert 'partial range "B2:C4" is not supported yet' in str(excinfo.value)


class TestLookupContract:
    def test_find_headers_row_reports_the_row_and_each_patterns_column(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.find_headers_row(
            opened.handle, "Data", "A1:C5", ["Region", "Amount"]
        ) == (2, {"Region": "A", "Amount": "B"})

    def test_find_headers_row_returns_none_when_no_row_matches_every_pattern(
        self, opened: Opened
    ) -> None:
        assert (
            opened.primitives.find_headers_row(
                opened.handle, "Data", "A1:C5", ["Region", "Nonexistent"]
            )
            is None
        )

    def test_find_row_locates_a_value_by_equality(self, opened: Opened) -> None:
        assert opened.primitives.find_row(opened.handle, "Data", "A", "South") == 4

    def test_find_row_starts_after_the_header_row_when_given_one(
        self, opened: Opened
    ) -> None:
        assert (
            opened.primitives.find_row(
                opened.handle, "Data", "A", "Region", header_row=2
            )
            is None
        )

    def test_find_row_returns_none_when_the_value_is_absent(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.find_row(opened.handle, "Data", "A", "Nowhere") is None

    def test_find_column_matches_a_header_by_regex(self, opened: Opened) -> None:
        assert (
            opened.primitives.find_column(opened.handle, "Data", 2, "^Amount$") == "B"
        )

    def test_find_column_returns_none_when_no_header_matches(
        self, opened: Opened
    ) -> None:
        assert (
            opened.primitives.find_column(opened.handle, "Data", 2, "Missing") is None
        )

    def test_find_columns_omits_patterns_that_matched_nothing(
        self, opened: Opened
    ) -> None:
        assert opened.primitives.find_columns(
            opened.handle, "Data", 2, {"amt": "Amount", "gone": "Missing"}
        ) == {"amt": "B"}
