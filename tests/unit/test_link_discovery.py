"""Unit tests for the static, pure-Python external-link discovery pass in excel_runner.engine
(docs/recalc_and_link_refresh_plan.md sec 2): classifying and resolving link Target strings
requires no Excel/COM at all (real-Excel is only needed to *create* fixtures for
`scan_external_link_targets`, whose own implementation is pure zipfile/XML).
"""

import zipfile
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

from excel_runner.backends import OwnedInstanceRegistry
from excel_runner.core import WorkbookRef
from excel_runner.engine import (
    classify_link_target,
    discover_write_intent_link_graph,
    inspect_save_blockers,
    inspection_paths,
    resolve_link_target,
    scan_external_link_targets,
)
from tests.unit.conftest import (
    plain_workbook,
    requires_excel,
    requires_working_xlwings_save,
    workbook_with_external_link,
)


class TestInspectionPaths:
    """Which file actually gets read when a workbook is inspected before the run starts
    (docs/backend_eligibility_build_plan.md W2). A workbook that doesn't exist yet still has
    inspectable content — its template's — and that content is what the run will inherit.
    """

    def test_existing_workbook_is_inspected_as_itself(self, tmp_path: Path) -> None:
        real = workbook_with_external_link(tmp_path / "manip.xlsx")
        workbooks = {"manip": WorkbookRef(name="manip", file=str(real))}

        assert inspection_paths(workbooks) == {"manip": real}

    def test_missing_workbook_is_inspected_through_its_template(
        self, tmp_path: Path
    ) -> None:
        template = workbook_with_external_link(tmp_path / "template.xlsx")
        workbooks = {
            "template": WorkbookRef(name="template", file=str(template)),
            "report": WorkbookRef(
                name="report",
                file=str(tmp_path / "report.xlsx"),
                create_if_missing=True,
                template="template",
            ),
        }

        assert inspection_paths(workbooks) == {
            "template": template,
            "report": template,
        }

    def test_missing_workbook_with_no_template_is_simply_absent(
        self, tmp_path: Path
    ) -> None:
        """Nothing to read, so nothing to say about it — not an error at this stage."""
        workbooks = {
            "blank": WorkbookRef(
                name="blank", file=str(tmp_path / "blank.xlsx"), create_if_missing=True
            )
        }

        assert inspection_paths(workbooks) == {}

    def test_missing_workbook_whose_template_also_does_not_exist_is_absent(
        self, tmp_path: Path
    ) -> None:
        workbooks = {
            "template": WorkbookRef(name="template", file=str(tmp_path / "gone.xlsx")),
            "report": WorkbookRef(
                name="report",
                file=str(tmp_path / "report.xlsx"),
                create_if_missing=True,
                template="template",
            ),
        }

        assert inspection_paths(workbooks) == {}


class TestDiscoverLinkGraphThroughTemplates:
    """A `create_if_missing` workbook inherits its template's links the moment it is created,
    so the link graph must see them on the *first* run — not only on the second, once the file
    happens to exist (docs/backend_eligibility_build_plan.md W2).
    """

    def test_link_carried_by_a_template_is_discovered_before_the_file_exists(
        self, tmp_path: Path
    ) -> None:
        target = plain_workbook(tmp_path / "target.xlsx")
        template = workbook_with_external_link(
            tmp_path / "template.xlsx", target=str(target.resolve())
        )
        report = tmp_path / "report.xlsx"
        workbook_paths = {"report": report, "target": target}

        graph = discover_write_intent_link_graph(
            workbook_paths,
            write_intent={"report", "target"},
            scan_paths={"report": template, "target": target},
        )

        assert graph == {"report": {"target"}, "target": set()}

    def test_without_scan_paths_a_not_yet_created_workbook_contributes_nothing(
        self, tmp_path: Path
    ) -> None:
        """The regression this fixes: no file on disk, so nothing was ever scanned, so the
        workbook looked link-free and got committed in arbitrary order."""
        target = plain_workbook(tmp_path / "target.xlsx")
        workbook_paths = {"report": tmp_path / "report.xlsx", "target": target}

        graph = discover_write_intent_link_graph(
            workbook_paths, write_intent={"report", "target"}
        )

        assert graph == {"report": set(), "target": set()}

    def test_a_templates_same_folder_link_is_still_not_an_r4_concern(
        self, tmp_path: Path
    ) -> None:
        """Scanning through a template widens *which file* is read, not *which links matter*.
        An R1 (same-folder) link is left out of the commit graph exactly as it would be if the
        workbook already existed and carried it directly.
        """
        target = plain_workbook(tmp_path / "target.xlsx")
        template = workbook_with_external_link(
            tmp_path / "template.xlsx", target="target.xlsx"
        )
        workbook_paths = {"report": tmp_path / "report.xlsx", "target": target}

        graph = discover_write_intent_link_graph(
            workbook_paths,
            write_intent={"report", "target"},
            scan_paths={"report": template, "target": target},
        )

        assert graph == {"report": set(), "target": set()}


class TestClassifyLinkTarget:
    def test_bare_filename_is_same_folder(self) -> None:
        assert classify_link_target("target.xlsx") == "same_folder"

    def test_forward_slash_subpath_is_relative_subpath(self) -> None:
        assert classify_link_target("other/target.xlsx") == "relative_subpath"

    def test_parent_relative_subpath_is_relative_subpath(self) -> None:
        assert classify_link_target("../target.xlsx") == "relative_subpath"

    def test_windows_drive_absolute_path_is_absolute(self) -> None:
        assert classify_link_target(r"C:\data\target.xlsx") == "absolute"

    def test_unc_path_is_absolute(self) -> None:
        assert classify_link_target(r"\\server\share\target.xlsx") == "absolute"

    def test_file_uri_is_absolute(self) -> None:
        assert classify_link_target("file:///C:/data/target.xlsx") == "absolute"

    def test_driveless_rooted_path_is_absolute(self) -> None:
        # What real Excel actually writes for a genuine external link to a file on the *same
        # drive* as the linking workbook — drive letter omitted, but still a real absolute
        # link (R3/R4), not same-folder/relative — confirmed via a real Excel-authored fixture
        # (TestScanExternalLinkTargets.test_finds_a_same_drive_absolute_link_target below).
        assert (
            classify_link_target("/Dev/projects/excel-runner/demos/target.xlsx")
            == "absolute"
        )


class TestResolveLinkTarget:
    def test_same_folder_target_resolves_against_linking_workbooks_folder(
        self, tmp_path: Path
    ) -> None:
        linking_path = tmp_path / "linking.xlsx"

        resolved = resolve_link_target("target.xlsx", linking_path)

        assert resolved == (tmp_path / "target.xlsx").resolve()

    def test_relative_subpath_resolves_against_linking_workbooks_folder(
        self, tmp_path: Path
    ) -> None:
        linking_path = tmp_path / "linking.xlsx"

        resolved = resolve_link_target("other/target.xlsx", linking_path)

        assert resolved == (tmp_path / "other" / "target.xlsx").resolve()

    def test_windows_absolute_path_resolves_to_itself(self, tmp_path: Path) -> None:
        linking_path = tmp_path / "linking.xlsx"
        absolute = tmp_path / "elsewhere" / "target.xlsx"

        resolved = resolve_link_target(str(absolute), linking_path)

        assert resolved == absolute.resolve()

    def test_file_uri_resolves_to_the_plain_path(self, tmp_path: Path) -> None:
        linking_path = tmp_path / "linking.xlsx"
        absolute = tmp_path / "elsewhere" / "target.xlsx"

        resolved = resolve_link_target(absolute.as_uri(), linking_path)

        assert resolved == absolute.resolve()


@requires_excel
@requires_working_xlwings_save
class TestScanExternalLinkTargets:
    def test_finds_a_same_folder_link_target(self, tmp_path: Path) -> None:
        registry = OwnedInstanceRegistry()
        app = registry.spawn()
        try:
            target = app.books.add()
            target.sheets[0].range("A1").value = 5
            target.save(str(tmp_path / "target.xlsx"))
            target.close()

            linking = app.books.add()
            linking.sheets[0].range("A1").formula = "='[target.xlsx]Sheet1'!A1*2"
            linking.save(str(tmp_path / "linking.xlsx"))
            linking.close()

            targets = scan_external_link_targets(tmp_path / "linking.xlsx")

            assert targets == ["target.xlsx"]
        finally:
            registry.close_owned()

    def test_finds_a_subfolder_link_target(self, tmp_path: Path) -> None:
        (tmp_path / "other").mkdir()
        registry = OwnedInstanceRegistry()
        app = registry.spawn()
        try:
            target = app.books.add()
            target.sheets[0].range("A1").value = 7
            target.save(str(tmp_path / "other" / "target2.xlsx"))
            target.close()

            linking = app.books.add()
            linking.sheets[0].range(
                "A1"
            ).formula = f"='{tmp_path / 'other'}\\[target2.xlsx]Sheet1'!A1*2"
            linking.save(str(tmp_path / "linking.xlsx"))
            linking.close()

            targets = scan_external_link_targets(tmp_path / "linking.xlsx")

            assert targets == ["other/target2.xlsx"]
        finally:
            registry.close_owned()

    def test_finds_a_same_drive_absolute_link_target(self, tmp_path: Path) -> None:
        # Typing the absolute path directly (rather than same-folder + ChangeLink) is
        # deliberate: Excel silently re-normalizes a ChangeLink'd target back to same-folder
        # when both files happen to sit in the same folder at ChangeLink time.
        (tmp_path / "linking_dir").mkdir()
        registry = OwnedInstanceRegistry()
        app = registry.spawn()
        try:
            target = app.books.add()
            target.sheets[0].range("A1").value = 9
            target.save(str(tmp_path / "target.xlsx"))
            target.close()

            linking = app.books.add()
            linking.sheets[0].range(
                "A1"
            ).formula = f"='{tmp_path}\\[target.xlsx]Sheet1'!A1*2"
            linking.save(str(tmp_path / "linking_dir" / "linking.xlsx"))
            linking.close()

            targets = scan_external_link_targets(
                tmp_path / "linking_dir" / "linking.xlsx"
            )

            # Excel writes two Target relationships for one same-drive absolute link: a
            # relative fallback (ignored — classifies as relative_subpath, R2) and the real
            # drive-omitted-rooted absolute one (R3/R4) that discover_write_intent_link_graph
            # actually keys off of.
            absolute_targets = [
                t for t in targets if classify_link_target(t) == "absolute"
            ]
            assert len(absolute_targets) == 1
            assert (
                resolve_link_target(
                    absolute_targets[0], tmp_path / "linking_dir" / "linking.xlsx"
                )
                == (tmp_path / "target.xlsx").resolve()
            )
        finally:
            registry.close_owned()

    def test_returns_empty_list_with_no_external_links(self, tmp_path: Path) -> None:
        registry = OwnedInstanceRegistry()
        app = registry.spawn()
        try:
            book = app.books.add()
            book.sheets[0].range("A1").value = "no links here"
            book.save(str(tmp_path / "standalone.xlsx"))
            book.close()

            assert scan_external_link_targets(tmp_path / "standalone.xlsx") == []
        finally:
            registry.close_owned()


class TestInternalReferencesAreNotExternalLinks:
    """A false positive here is expensive, not just untidy: one phantom target makes
    `inspect_save_blockers` report OUTBOUND_EXTERNAL_LINKS, which promotes the workbook off
    openpyxl onto a live Excel session for the whole run
    (docs/backend_eligibility_build_plan.md sec 1.3, and promotion is sticky). A workbook
    whose formulas merely reference its own other sheets must stay on the fast path.

    Built with openpyxl rather than live Excel deliberately: these assert what the scanner
    does *not* pick up, so they need to run everywhere, and the on-disk shapes involved
    (formula text in the sheet XML, a definedName, a hyperlink relationship) are ordinary
    OOXML that doesn't need Excel to author.
    """

    def _sheet_rels(self, path: Path) -> list[str]:
        with zipfile.ZipFile(path) as archive:
            return [
                name
                for name in archive.namelist()
                if name.startswith("xl/worksheets/_rels/")
            ]

    def _cross_sheet_workbook(self, path: Path) -> Path:
        workbook = openpyxl.Workbook()
        data = workbook.active
        assert data is not None
        data.title = "Data"
        data["A1"] = 5
        summary = workbook.create_sheet("Summary")
        summary["A1"] = "=Data!A1*2"
        summary["A2"] = "=SUM(Data!A1:A10)"
        summary["A3"] = "='Data'!A1"
        workbook.save(path)
        return path

    def test_a_cross_sheet_formula_is_not_an_external_link(self, tmp_path: Path) -> None:
        path = self._cross_sheet_workbook(tmp_path / "internal.xlsx")

        assert scan_external_link_targets(path) == []

    def test_a_cross_sheet_formula_does_not_block_the_file_backend(
        self, tmp_path: Path
    ) -> None:
        """The consequence that actually matters — this is what decides whether the workbook
        keeps the openpyxl fast path or gets promoted to a live Excel session."""
        path = self._cross_sheet_workbook(tmp_path / "internal.xlsx")

        assert inspect_save_blockers(path) == frozenset()

    def test_a_defined_name_spanning_sheets_is_not_an_external_link(
        self, tmp_path: Path
    ) -> None:
        workbook = openpyxl.Workbook()
        data = workbook.active
        assert data is not None
        data.title = "Data"
        data["A1"] = 5
        workbook.defined_names.add(DefinedName("Total", attr_text="Data!$A$1"))
        path = tmp_path / "named.xlsx"
        workbook.save(path)

        assert scan_external_link_targets(path) == []
        assert inspect_save_blockers(path) == frozenset()

    def test_a_hyperlink_is_not_an_external_link(self, tmp_path: Path) -> None:
        """The sharpest case: a hyperlink really is stored as a `TargetMode="External"`
        relationship, so the `TargetMode` check alone would happily pick it up. What keeps it
        out is that the scanner only reads `xl/externalLinks/_rels/`, and a hyperlink lives in
        `xl/worksheets/_rels/`. A workbook is not link-bearing just because a cell points at
        a website."""
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet["A1"] = "see docs"
        sheet["A1"].hyperlink = "https://example.com/reference.xlsx"
        path = tmp_path / "hyperlinked.xlsx"
        workbook.save(path)

        # Guard: if openpyxl stopped writing the hyperlink as an external relationship this
        # test would pass for the wrong reason, proving nothing about the scanner.
        assert self._sheet_rels(path) != []
        with zipfile.ZipFile(path) as archive:
            assert b'TargetMode="External"' in archive.read(self._sheet_rels(path)[0])

        assert scan_external_link_targets(path) == []
        assert inspect_save_blockers(path) == frozenset()

    def test_a_genuine_external_link_is_still_detected(self, tmp_path: Path) -> None:
        """The other half of the contract: tightening against internal references must not
        have made the scanner blind."""
        path = workbook_with_external_link(tmp_path / "linked.xlsx", "target.xlsx")

        assert scan_external_link_targets(path) == ["target.xlsx"]
